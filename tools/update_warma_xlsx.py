#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
自动更新 Warma 相关 / warma养鸽场 相关 两个 Excel 表。

用法:
  python tools/update_warma_xlsx.py            # 读取缓存 + 只补充缺失的 B 站元数据
  python tools/update_warma_xlsx.py --refresh  # 强制重新拉取 B 站元数据
  python tools/update_warma_xlsx.py --skip-fetch # 完全离线，使用当前缓存

核心数据源:
  registry.json                 # 字幕库整理出来的视频清单
  @Warma 相关.xlsx             # 历史主号表，保留其中的时长与未匹配记录
  @warma养鸽场 相关.xlsx       # 历史小号表，保留其中的时长与未匹配记录
"""

from __future__ import annotations

import argparse
import datetime as dt
import gzip
import json
import re
import shutil
import sys
import time
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# 项目目录（registry.json 所在位置）
PROJECT_DIR = Path(__file__).resolve().parents[1]
# Excel 文件所在目录（比项目目录高一级，即 F:\warma百科）
XLSX_DIR = Path(__file__).resolve().parents[2]
MAIN_FILE = XLSX_DIR / "@Warma 相关.xlsx"
SMALL_FILE = XLSX_DIR / "@warma养鸽场 相关.xlsx"
REGISTRY_FILE = PROJECT_DIR / "registry.json"
CACHE_FILE = PROJECT_DIR / "tools" / "bili_meta_cache.json"
BACKUP_DIR = XLSX_DIR / "backup_xlsx"

CN_WEEK = ["周一", "周二", "周三", "周四", "周五", "周六", "周日"]
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
SUBHEAD_FILL = PatternFill("solid", fgColor="E0E7FF")
WARN_FILL = PatternFill("solid", fgColor="FEF3C7")
GOOD_FILL = PatternFill("solid", fgColor="DCFCE7")
BAD_FILL = PatternFill("solid", fgColor="FEE2E2")
CARD_FILL = PatternFill("solid", fgColor="EEF2FF")
THIN_GRAY = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)


def log(message: str) -> None:
    print(message, flush=True)


def normalize_title(value) -> str:
    return re.sub(r"[^0-9A-Za-z\u4e00-\u9fff#]+", "", str(value or "")).lower()


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def excel_time_to_seconds(value) -> float | int | None:
    if value in (None, ""):
        return None
    if isinstance(value, dt.time):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, dt.datetime):
        return value.hour * 3600 + value.minute * 60 + value.second
    if isinstance(value, (int, float)):
        # openpyxl 有时会把 Excel 时间读成天数。
        return round(float(value) * 86400) if 0 < float(value) < 2 else float(value)
    text = str(value).strip()
    parts = re.findall(r"\d+", text)
    if not parts:
        return None
    nums = [int(x) for x in parts]
    if len(nums) >= 3:
        return nums[-3] * 3600 + nums[-2] * 60 + nums[-1]
    if len(nums) == 2:
        return nums[0] * 60 + nums[1]
    return nums[0]


def read_history_rows(path: Path, max_rows: int = 1200):
    """只读取旧表第一页的核心列，避免读取异常膨胀的工作表维度。"""
    rows = []
    if not path.exists():
        return rows
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, max_rows), max_col=13, values_only=True):
        title = row[1]
        if not title:
            continue
        link = row[11] or row[6] or ""
        match = re.search(r"(BV[0-9A-Za-z]{10})", str(link))
        rows.append(
            {
                "old_no": row[0],
                "title": str(title).strip(),
                "date": row[3],
                "duration_seconds": excel_time_to_seconds(row[4]),
                "link": str(link) if link else None,
                "bvid": match.group(1) if match else None,
            }
        )
    wb.close()
    return rows


def fetch_bili_meta(bvid: str, cookie: str | None = None):
    url = "https://api.bilibili.com/x/web-interface/view?" + urllib.parse.urlencode({"bvid": bvid})
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/126 Safari/537.36",
        "Referer": "https://www.bilibili.com/",
        "Accept": "application/json,text/plain,*/*",
    }
    if cookie:
        headers["Cookie"] = cookie
    for attempt in range(3):
        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=12) as response:
                raw = response.read()
                if response.headers.get("Content-Encoding") == "gzip":
                    raw = gzip.decompress(raw)
                payload = json.loads(raw.decode("utf-8"))
            if payload.get("code") == 0:
                return payload.get("data") or {}
            if payload.get("code") in (-400, -404, 62002, 62012):
                return {"_error": payload.get("message", "not found"), "_code": payload.get("code")}
        except Exception as exc:
            if attempt == 2:
                return {"_error": str(exc)}
        time.sleep(0.4 + attempt * 0.7)
    return {"_error": "request failed"}


def refresh_meta_cache(bvids: list[str], force: bool, skip_fetch: bool) -> dict:
    cache = read_json(CACHE_FILE) if CACHE_FILE.exists() else {}
    if skip_fetch:
        missing = [bv for bv in bvids if bv not in cache]
        log(f"离线模式：缺少 {len(missing)} 条 B 站元数据")
        return cache
    todo = [bv for bv in bvids if force or bv not in cache]
    log(f"B 站元数据：缓存 {len(cache) - len([b for b in todo if b in cache])} 条，本次请求 {len(todo)} 条")
    cookie = None
    for index, bvid in enumerate(todo, 1):
        cache[bvid] = fetch_bili_meta(bvid, cookie)
        if index % 25 == 0 or index == len(todo):
            log(f"  已拉取 {index}/{len(todo)}")
        if index % 100 == 0:
            write_json(CACHE_FILE, cache)
        time.sleep(0.14)
    write_json(CACHE_FILE, cache)
    return cache


def build_unified_rows(main_rows: list[dict], small_rows: list[dict], registry_rows: list[dict], meta_cache: dict):
    reg_by_bvid = {row["bvid"]: row for row in registry_rows}
    reg_by_title = {normalize_title(row["title"]): row for row in registry_rows}
    used_registry_bvids = set()

    def api_date(bvid, fallback):
        info = meta_cache.get(bvid) if bvid else None
        pubdate = info.get("pubdate") if isinstance(info, dict) else None
        if pubdate:
            return dt.datetime.fromtimestamp(int(pubdate), ZoneInfo("Asia/Shanghai")).replace(tzinfo=None)
        return fallback

    def make_record(source_row, workbook_account, source_label):
        bvid = source_row.get("bvid")
        reg = reg_by_bvid.get(bvid) or reg_by_title.get(normalize_title(source_row.get("title")))
        if reg:
            bvid = reg["bvid"]
            used_registry_bvids.add(bvid)
        info = meta_cache.get(bvid) if bvid else None
        if not isinstance(info, dict):
            info = {}
        title = info.get("title") or (reg or {}).get("title") or source_row.get("title")
        date = api_date(bvid, source_row.get("date"))
        duration = info.get("duration") or source_row.get("duration_seconds")
        stat = info.get("stat") or {}
        return {
            "title": title,
            "bvid": bvid,
            "aid": info.get("aid") or (str(source_row.get("old_no")) if False else None),
            "date": date,
            "year": date.year if isinstance(date, dt.datetime) else None,
            "month": date.month if isinstance(date, dt.datetime) else None,
            "type": (reg or {}).get("type") or info.get("tname") or None,
            "account": (reg or {}).get("account") or workbook_account,
            "duration_seconds": int(duration) if duration not in (None, "") else None,
            "sub_lines": (reg or {}).get("sub_lines"),
            "sub_chars": (reg or {}).get("sub_chars"),
            "view": stat.get("view"),
            "danmaku": stat.get("danmaku"),
            "reply": stat.get("reply"),
            "favorite": stat.get("favorite"),
            "coin": stat.get("coin"),
            "share": stat.get("share"),
            "like": stat.get("like"),
            "link": f"https://www.bilibili.com/video/{bvid}" if bvid else source_row.get("link"),
            "source": source_label + (" + 字幕库" if reg else " + B站API" if bvid else " + 历史表格"),
            "note": "历史表保留记录" if not reg else None,
        }

    unified = []
    unmatched_main = []
    unmatched_small = []

    for row in main_rows:
        matched_bvid = row.get("bvid")
        matched = bool(matched_bvid and matched_bvid in reg_by_bvid) or normalize_title(row.get("title")) in reg_by_title
        record = make_record(row, "主号（Warma）", "历史主号表")
        unified.append(record)
        if not matched:
            unmatched_main.append(record)

    for row in small_rows:
        matched_bvid = row.get("bvid")
        matched = bool(matched_bvid and matched_bvid in reg_by_bvid) or normalize_title(row.get("title")) in reg_by_title
        record = make_record(row, "小号（warma养鸽场）", "历史小号表")
        unified.append(record)
        if not matched:
            unmatched_small.append(record)

    existing_keys = {
        (normalize_title(row["title"]), str(row["date"])[:10]) for row in unified if row.get("title") and row.get("date")
    }
    for row in registry_rows:
        if row["bvid"] in used_registry_bvids:
            continue
        date = api_date(row["bvid"], dt.datetime.fromisoformat(row["date"]))
        key = (normalize_title(row["title"]), str(date)[:10])
        if key in existing_keys:
            continue
        info = meta_cache.get(row["bvid"]) or {}
        stat = info.get("stat") or {}
        unified.append(
            {
                "title": info.get("title") or row["title"],
                "bvid": row["bvid"],
                "aid": info.get("aid"),
                "date": date,
                "year": date.year,
                "month": date.month,
                "type": row.get("type") or info.get("tname"),
                "account": row.get("account"),
                "duration_seconds": int(info.get("duration")) if info.get("duration") else None,
                "sub_lines": row.get("sub_lines"),
                "sub_chars": row.get("sub_chars"),
                "view": stat.get("view"),
                "danmaku": stat.get("danmaku"),
                "reply": stat.get("reply"),
                "favorite": stat.get("favorite"),
                "coin": stat.get("coin"),
                "share": stat.get("share"),
                "like": stat.get("like"),
                "link": f"https://www.bilibili.com/video/{row['bvid']}",
                "source": "字幕库 + B站API",
                "note": None,
            }
        )

    # 按 BV 去重，再按标题+日期去重。
    seen_bv = set()
    seen_key = set()
    deduped = []
    for row in sorted(unified, key=lambda x: (x["date"] or dt.datetime(1970, 1, 1), x["bvid"] or "", x["title"] or "")):
        if row["bvid"] and row["bvid"] in seen_bv:
            continue
        key = (normalize_title(row["title"]), str(row["date"])[:10])
        if key in seen_key:
            continue
        if row["bvid"]:
            seen_bv.add(row["bvid"])
        seen_key.add(key)
        deduped.append(row)
    return deduped, unmatched_main, unmatched_small


def style_header(ws, row, last_col):
    for col in range(1, last_col + 1):
        cell = ws.cell(row, col)
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 26


def set_widths(ws, widths: list[float]):
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def add_note(ws, row, title, text, fill=None, width=8):
    ws.cell(row, 1, title).font = Font(bold=True, size=12, color="1F2937")
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=width)
    cell = ws.cell(row, 2, text)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    if fill:
        ws.cell(row, 1).fill = fill
        cell.fill = fill


def write_data_sheet(wb: Workbook, rows: list[dict], extra_formula_rows: int = 200):
    # 删除 openpyxl 默认创建的空 Sheet，避免 wb.active 指向已合并单元格的说明页。
    for name in list(wb.sheetnames):
        sheet = wb[name]
        if (
            sheet.max_row == 1
            and sheet.max_column == 1
            and sheet["A1"].value is None
        ):
            wb.remove(sheet)
            break
    ws = wb.create_sheet("数据总表")
    ws.title = "数据总表"
    headers = [
        "序号", "标题", "BV号", "AV号", "投稿日期", "年", "月", "星期", "内容类型", "账号",
        "时长(秒)", "时长", "距上次投稿(天)", "间隔等级", "字幕行数", "字幕字符数",
        "播放量", "弹幕", "评论", "收藏", "投币", "分享", "点赞", "点赞率", "互动率",
        "视频链接", "数据来源", "备注",
    ]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    for index, row in enumerate(rows, 1):
        r = index + 1
        date = row["date"]
        ws.cell(r, 1, index)
        ws.cell(r, 2, row["title"])
        ws.cell(r, 3, row["bvid"])
        ws.cell(r, 4, row["aid"])
        ws.cell(r, 5, date)
        ws.cell(r, 6, f'=IF($E{r}="","",YEAR($E{r}))')
        ws.cell(r, 7, f'=IF($E{r}="","",MONTH($E{r}))')
        ws.cell(r, 8, f'=IF($E{r}="","",CHOOSE(WEEKDAY($E{r},2),"周一","周二","周三","周四","周五","周六","周日"))')
        ws.cell(r, 9, row["type"])
        ws.cell(r, 10, row["account"])
        ws.cell(r, 11, row["duration_seconds"])
        ws.cell(r, 12, f'=IF($K{r}="","",TEXT($K{r}/86400,"[h]:mm:ss"))')
        ws.cell(r, 13, f'=IF(OR($B{r}="",$E{r}=""),"",$E{r}-$E{r-1})')
        ws.cell(r, 14, f'=IF($M{r}="","",IF($M{r}<=7,"≤7天",IF($M{r}<=30,"≤30天",IF($M{r}<=90,"≤90天",">90天"))))')
        ws.cell(r, 15, row["sub_lines"])
        ws.cell(r, 16, row["sub_chars"])
        ws.cell(r, 17, row["view"])
        ws.cell(r, 18, row["danmaku"])
        ws.cell(r, 19, row["reply"])
        ws.cell(r, 20, row["favorite"])
        ws.cell(r, 21, row["coin"])
        ws.cell(r, 22, row["share"])
        ws.cell(r, 23, row["like"])
        ws.cell(r, 24, f'=IFERROR($W{r}/$Q{r},"")')
        ws.cell(r, 25, f'=IFERROR(($W{r}+$T{r}+$U{r}+$V{r}+$R{r}+$S{r})/$Q{r},"")')
        ws.cell(r, 26, row["link"])
        ws.cell(r, 27, row["source"])
        ws.cell(r, 28, row["note"])

    # 预留公式行，方便手动追加新视频后自动计算。
    last_written = len(rows) + 1
    for r in range(last_written + 1, last_written + extra_formula_rows + 1):
        ws.cell(r, 6, f'=IF($E{r}="","",YEAR($E{r}))')
        ws.cell(r, 7, f'=IF($E{r}="","",MONTH($E{r}))')
        ws.cell(r, 8, f'=IF($E{r}="","",CHOOSE(WEEKDAY($E{r},2),"周一","周二","周三","周四","周五","周六","周日"))')
        ws.cell(r, 12, f'=IF($K{r}="","",TEXT($K{r}/86400,"[h]:mm:ss"))')
        ws.cell(r, 13, f'=IF(OR($B{r}="",$E{r}=""),"",$E{r}-$E{r-1})')
        ws.cell(r, 14, f'=IF($M{r}="","",IF($M{r}<=7,"≤7天",IF($M{r}<=30,"≤30天",IF($M{r}<=90,"≤90天",">90天"))))')
        ws.cell(r, 24, f'=IFERROR($W{r}/$Q{r},"")')
        ws.cell(r, 25, f'=IFERROR(($W{r}+$T{r}+$U{r}+$V{r}+$R{r}+$S{r})/$Q{r},"")')

    last_row = last_written + extra_formula_rows
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:AB{last_row}"
    ws.conditional_formatting.add(
        f"M2:M{last_row}",
        CellIsRule(operator="greaterThan", formula=["90"], fill=BAD_FILL, font=Font(color="B91C1C", bold=True)),
    )
    ws.conditional_formatting.add(
        f"K2:K{last_row}",
        FormulaRule(formula=[f'AND($B2<>"",$K2="")'], fill=WARN_FILL),
    )
    ws.conditional_formatting.add(
        f"O2:O{last_row}",
        FormulaRule(formula=[f'AND($B2<>"",$O2>=1000)'], fill=GOOD_FILL),
    )
    set_widths(
        ws,
        [6, 55, 15, 12, 12, 7, 7, 9, 16, 18, 10, 12, 16, 12, 10, 12, 12, 10, 10, 10, 10, 10, 10, 10, 10, 42, 22, 22],
    )
    for r in range(2, last_written + 1):
        ws.cell(r, 5).number_format = "yyyy-mm-dd"
        ws.cell(r, 11).number_format = "#,##0"
        ws.cell(r, 15).number_format = "#,##0"
        ws.cell(r, 16).number_format = "#,##0"
        for c in range(17, 24):
            ws.cell(r, c).number_format = "#,##0"
        ws.cell(r, 24).number_format = "0.00%"
        ws.cell(r, 25).number_format = "0.00%"
        ws.cell(r, 2).alignment = Alignment(vertical="center", wrap_text=False)
        ws.row_dimensions[r].height = 18
    return ws


def write_dashboard(wb: Workbook, data_rows: int, years: list[int]):
    ws = wb.create_sheet("仪表盘")
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Warma 数据仪表盘"
    ws["A1"].font = Font(size=20, bold=True, color="4F46E5")
    ws["A2"] = "以下公式会随“数据总表”自动刷新；筛选数据后请查看年度分析表。"
    ws["A2"].font = Font(size=10, color="6B7280")
    kpis = [
        ("视频总数", f"=COUNTA('数据总表'!$B$2:$B${data_rows + 201})"),
        ("字幕行数", f"=SUM('数据总表'!$O$2:$O${data_rows + 201})"),
        ("字幕字符数", f"=SUM('数据总表'!$P$2:$P${data_rows + 201})"),
        ("总时长(小时)", f"=ROUND(SUM('数据总表'!$K$2:$K${data_rows + 201})/3600,1)"),
        ("平均时长(分)", f"=ROUND(AVERAGE('数据总表'!$K$2:$K${data_rows + 201})/60,1)"),
        ("平均间隔(天)", f"=ROUND(AVERAGE('数据总表'!$M$3:$M${data_rows + 201}),1)"),
        ("最新日期", f"=IFERROR(TEXT(MAX('数据总表'!$E$2:$E${data_rows + 201}),\"yyyy-mm-dd\"),\"-\")"),
        ("播放量合计", f"=SUM('数据总表'!$Q$2:$Q${data_rows + 201})"),
        ("点赞量合计", f"=SUM('数据总表'!$W$2:$W${data_rows + 201})"),
        ("平均互动率", f"=IFERROR(AVERAGE('数据总表'!$Y$2:$Y${data_rows + 201}),\"-\")"),
    ]
    for index, (label, formula) in enumerate(kpis):
        col = index % 5
        row = 4 + (index // 5) * 3
        ws.cell(row, 1 + col * 2, label).font = Font(size=10, color="6B7280")
        ws.cell(row + 1, 1 + col * 2, formula).font = Font(size=18, bold=True, color="111827")
        for rr in (row, row + 1):
            for cc in range(1 + col * 2, 3 + col * 2):
                ws.cell(rr, cc).fill = CARD_FILL
                ws.cell(rr, cc).border = BORDER
        ws.merge_cells(start_row=row, start_column=1 + col * 2, end_row=row, end_column=2 + col * 2)
        ws.merge_cells(start_row=row + 1, start_column=1 + col * 2, end_row=row + 1, end_column=2 + col * 2)

    # 图表数据来自年度分析和类型分析。
    chart1 = BarChart()
    chart1.title = "年度投稿数"
    chart1.height, chart1.width = 9, 18
    chart1.y_axis.title = "视频数"
    chart1.add_data(Reference(wb["年度分析"], min_col=2, min_row=1, max_row=1 + len(years)), titles_from_data=True)
    chart1.set_categories(Reference(wb["年度分析"], min_col=1, min_row=2, max_row=1 + len(years)))
    ws.add_chart(chart1, "A11")

    chart2 = LineChart()
    chart2.title = "年度字幕行数"
    chart2.height, chart2.width = 9, 18
    chart2.add_data(Reference(wb["年度分析"], min_col=5, min_row=1, max_row=1 + len(years)), titles_from_data=True)
    chart2.set_categories(Reference(wb["年度分析"], min_col=1, min_row=2, max_row=1 + len(years)))
    ws.add_chart(chart2, "K11")

    chart3 = PieChart()
    chart3.title = "内容类型分布"
    chart3.height, chart3.width = 10, 18
    type_count = min(20, wb["类型分析"].max_row - 1)
    chart3.add_data(Reference(wb["类型分析"], min_col=2, min_row=1, max_row=1 + type_count), titles_from_data=True)
    chart3.set_categories(Reference(wb["类型分析"], min_col=1, min_row=2, max_row=1 + type_count))
    ws.add_chart(chart3, "A31")

    chart4 = BarChart()
    chart4.title = "视频时长分段"
    chart4.height, chart4.width = 10, 18
    bucket_count = wb["时长分段"].max_row - 1
    chart4.add_data(Reference(wb["时长分段"], min_col=2, min_row=1, max_row=1 + bucket_count), titles_from_data=True)
    chart4.set_categories(Reference(wb["时长分段"], min_col=1, min_row=2, max_row=1 + bucket_count))
    ws.add_chart(chart4, "K31")
    set_widths(ws, [18] * 10)


def write_year_analysis(wb: Workbook, rows: list[dict], data_rows: int):
    ws = wb.create_sheet("年度分析")
    years = sorted({row["date"].year for row in rows if row["date"]})
    headers = ["年份", "视频数", "总时长(秒)", "平均时长(秒)", "字幕行数", "字幕字符数", "平均间隔(天)", "播放量", "点赞量"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for year in years:
        r = ws.max_row + 1
        ws.cell(r, 1, year)
        ws.cell(r, 2, f"=COUNTIFS('数据总表'!$F$2:$F${data_rows + 201},$A{r})")
        ws.cell(r, 3, f"=SUMIFS('数据总表'!$K$2:$K${data_rows + 201},'数据总表'!$F$2:$F${data_rows + 201},$A{r})")
        ws.cell(r, 4, f"=IFERROR(AVERAGEIFS('数据总表'!$K$2:$K${data_rows + 201},'数据总表'!$F$2:$F${data_rows + 201},$A{r}),\"\")")
        ws.cell(r, 5, f"=SUMIFS('数据总表'!$O$2:$O${data_rows + 201},'数据总表'!$F$2:$F${data_rows + 201},$A{r})")
        ws.cell(r, 6, f"=SUMIFS('数据总表'!$P$2:$P${data_rows + 201},'数据总表'!$F$2:$F${data_rows + 201},$A{r})")
        ws.cell(r, 7, f"=IFERROR(AVERAGEIFS('数据总表'!$M$2:$M${data_rows + 201},'数据总表'!$F$2:$F${data_rows + 201},$A{r}),\"\")")
        ws.cell(r, 8, f"=SUMIFS('数据总表'!$Q$2:$Q${data_rows + 201},'数据总表'!$F$2:$F${data_rows + 201},$A{r})")
        ws.cell(r, 9, f"=SUMIFS('数据总表'!$W$2:$W${data_rows + 201},'数据总表'!$F$2:$F${data_rows + 201},$A{r})")
    r = ws.max_row + 1
    ws.cell(r, 1, "合计")
    for col in (2, 3, 5, 6, 8, 9):
        ws.cell(r, col, f"=SUM({get_column_letter(col)}2:{get_column_letter(col)}{r - 1})")
    ws.cell(r, 4, f"=IFERROR(AVERAGE(D2:D{r - 1}),\"\")")
    ws.cell(r, 7, f"=IFERROR(AVERAGE(G2:G{r - 1}),\"\")")
    for cell in ws[r]:
        cell.fill = SUBHEAD_FILL
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
            if cell.column >= 2:
                cell.number_format = "#,##0"
    ws.freeze_panes = "A2"
    set_widths(ws, [10, 10, 14, 14, 12, 14, 14, 14, 14])
    return years


def write_type_analysis(wb: Workbook, rows: list[dict], data_rows: int):
    ws = wb.create_sheet("类型分析")
    type_counts = Counter(row["type"] for row in rows if row["type"])
    headers = ["内容类型", "视频数", "总时长(秒)", "平均时长(秒)", "字幕行数", "字幕字符数", "平均播放量", "平均互动率"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for type_name, _count in type_counts.most_common():
        r = ws.max_row + 1
        ws.cell(r, 1, type_name)
        ws.cell(r, 2, f"=COUNTIFS('数据总表'!$I$2:$I${data_rows + 201},$A{r})")
        ws.cell(r, 3, f"=SUMIFS('数据总表'!$K$2:$K${data_rows + 201},'数据总表'!$I$2:$I${data_rows + 201},$A{r})")
        ws.cell(r, 4, f"=IFERROR(AVERAGEIFS('数据总表'!$K$2:$K${data_rows + 201},'数据总表'!$I$2:$I${data_rows + 201},$A{r}),\"\")")
        ws.cell(r, 5, f"=SUMIFS('数据总表'!$O$2:$O${data_rows + 201},'数据总表'!$I$2:$I${data_rows + 201},$A{r})")
        ws.cell(r, 6, f"=SUMIFS('数据总表'!$P$2:$P${data_rows + 201},'数据总表'!$I$2:$I${data_rows + 201},$A{r})")
        ws.cell(r, 7, f"=IFERROR(AVERAGEIFS('数据总表'!$Q$2:$Q${data_rows + 201},'数据总表'!$I$2:$I${data_rows + 201},$A{r}),\"\")")
        ws.cell(r, 8, f"=IFERROR(AVERAGEIFS('数据总表'!$Y$2:$Y${data_rows + 201},'数据总表'!$I$2:$I${data_rows + 201},$A{r}),\"\")")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
            if cell.column in (2, 3, 5, 6, 7):
                cell.number_format = "#,##0"
            if cell.column == 8:
                cell.number_format = "0.00%"
    ws.freeze_panes = "A2"
    set_widths(ws, [20, 10, 14, 14, 12, 14, 14, 14])


def write_duration_buckets(wb: Workbook, data_rows: int):
    ws = wb.create_sheet("时长分段")
    buckets = [
        ("0-1分钟", 0, 59), ("1-2分钟", 60, 119), ("2-5分钟", 120, 299), ("5-10分钟", 300, 599),
        ("10-30分钟", 600, 1799), ("30-60分钟", 1800, 3599), ("1-2小时", 3600, 7199),
        ("2小时以上", 7200, 10**9),
    ]
    ws.append(["时长分段", "视频数"])
    style_header(ws, 1, 2)
    for name, low, high in buckets:
        r = ws.max_row + 1
        ws.cell(r, 1, name)
        ws.cell(r, 2, f"=COUNTIFS('数据总表'!$K$2:$K${data_rows + 201},\">={low}\",'数据总表'!$K$2:$K${data_rows + 201},\"<={high}\")")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
        for cell in row:
            cell.border = BORDER
            if cell.column == 2:
                cell.number_format = "#,##0"
    set_widths(ws, [18, 12])


def write_unmatched(wb: Workbook, unmatched: list[dict], label: str):
    ws = wb.create_sheet("未匹配记录")
    ws["A1"] = "以下记录来自旧表，但未能在当前 registry.json 中精确匹配，已完整保留。"
    ws["A1"].font = Font(size=12, bold=True)
    headers = ["旧表序号", "标题", "日期", "时长(秒)", "链接", "来源"]
    ws.append([])
    ws.append(headers)
    style_header(ws, 3, len(headers))
    for row in unmatched:
        ws.append([row.get("old_no"), row.get("title"), row.get("date"), row.get("duration_seconds"), row.get("link"), label])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=len(headers)):
        for cell in row:
            cell.border = BORDER
            if cell.column == 4:
                cell.number_format = "#,##0"
    ws.freeze_panes = "A4"
    set_widths(ws, [12, 70, 14, 12, 55, 18])


def write_quality(wb: Workbook, data_rows: int):
    ws = wb.create_sheet("数据质量")
    ws["A1"] = "数据质量检查"
    ws["A1"].font = Font(size=16, bold=True, color="4F46E5")
    checks = [
        ("视频总数", f"=COUNTA('数据总表'!$B$2:$B${data_rows + 201})"),
        ("缺少 BV 号", f"=COUNTBLANK('数据总表'!$C$2:$C${data_rows + 201})"),
        ("缺少 AV 号", f"=COUNTBLANK('数据总表'!$D$2:$D${data_rows + 201})"),
        ("缺少投稿日期", f"=COUNTBLANK('数据总表'!$E$2:$E${data_rows + 201})"),
        ("缺少时长", f"=COUNTBLANK('数据总表'!$K$2:$K${data_rows + 201})"),
        ("缺少内容类型", f"=COUNTBLANK('数据总表'!$I$2:$I${data_rows + 201})"),
        ("缺少字幕行数", f"=COUNTBLANK('数据总表'!$O$2:$O${data_rows + 201})"),
        ("缺少播放量", f"=COUNTBLANK('数据总表'!$Q$2:$Q${data_rows + 201})"),
        ("拖更间隔 > 90天", f"=COUNTIF('数据总表'!$M$2:$M${data_rows + 201},\">90\")"),
        ("字幕 >= 1000行", f"=COUNTIF('数据总表'!$O$2:$O${data_rows + 201},\">=1000\")"),
    ]
    ws.append([])
    ws.append(["检查项", "数量"])
    style_header(ws, 3, 2)
    for name, formula in checks:
        ws.append([name, formula])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=2):
        for cell in row:
            cell.border = BORDER
            if cell.column == 2:
                cell.number_format = "#,##0"
    set_widths(ws, [24, 14])


def write_guide(wb: Workbook):
    ws = wb.create_sheet("使用说明", 0)
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Warma Excel 数据库"
    ws["A1"].font = Font(size=24, bold=True, color="4F46E5")
    ws["A2"] = "已接入 registry.json 与 B 站公开元数据，支持公式分析和图表联动。"
    ws["A2"].font = Font(size=11, color="6B7280")
    notes = [
        ("数据总表", "每行一个视频。年/月/星期/时长/间隔/互动率等列都是公式，筛选后仍会保留完整计算结果。"),
        ("仪表盘", "核心 KPI 和 4 张图表，数据来自年度分析、类型分析、时长分段。"),
        ("年度分析", "按年份统计视频数、总时长、字幕规模、播放和点赞。"),
        ("类型分析", "按内容类型统计规模、播放和互动表现。"),
        ("时长分段", "把视频分成 8 个时长区间，便于看出长视频/短视频构成。"),
        ("未匹配记录", "旧表中有、但 registry.json 精确匹配不到的视频会被完整保留，不会丢失。"),
        ("数据质量", "检查 BV 号、日期、时长、播放量、字幕行数等缺失情况。"),
        ("更新日志", "每次自动更新都会记录时间、视频数和主要变化。"),
        ("一键更新", "以后双击项目目录里的“一键更新表格.bat”即可重新同步 registry.json 和 B 站元数据。"),
    ]
    row = 4
    for title, text in notes:
        ws.cell(row, 1, title).font = Font(size=12, bold=True, color="111827")
        ws.cell(row, 2, text).font = Font(size=11, color="374151")
        ws.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.cell(row, 1).fill = SUBHEAD_FILL
        ws.cell(row, 2).fill = PatternFill("solid", fgColor="F8FAFC")
        ws.row_dimensions[row].height = 30
        row += 1
    set_widths(ws, [18, 22, 22, 22, 22, 22, 22, 22])


def write_update_log(wb: Workbook, rows_count: int, unmatched_count: int, message: str):
    ws = wb.create_sheet("更新日志")
    ws.append(["时间", "视频数", "未匹配数", "更新说明"])
    style_header(ws, 1, 4)
    now = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws.append([now, rows_count, unmatched_count, message])
    ws.freeze_panes = "A2"
    set_widths(ws, [20, 12, 12, 80])


def build_workbook(
    rows: list[dict],
    unmatched: list[dict],
    label: str,
    log_message: str,
    output: Path,
    account_filter: str | None = None,
):
    """account_filter: 例如 "主号" 只保留主号视频，"小号" 保留所有小号。"""
    if account_filter:
        rows = [r for r in rows if (r.get("account") or "").startswith(account_filter)]
    wb = Workbook()
    write_guide(wb)
    data_ws = write_data_sheet(wb, rows)
    data_rows = len(rows)
    years = write_year_analysis(wb, rows, data_rows)
    write_type_analysis(wb, rows, data_rows)
    write_duration_buckets(wb, data_rows)
    write_dashboard(wb, data_rows, years)
    write_unmatched(wb, unmatched, label)
    write_quality(wb, data_rows)
    write_update_log(wb, len(rows), len(unmatched), log_message)
    # 把使用说明放第一页，数据总表第二页。
    wb.move_sheet("数据总表", offset=-(wb.sheetnames.index("数据总表") - 1))
    wb.save(output)


def backup_file(path: Path):
    if not path.exists():
        return
    BACKUP_DIR.mkdir(exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    target = BACKUP_DIR / f"{path.stem}-{stamp}{path.suffix}"
    shutil.copy2(path, target)
    log(f"已备份：{target}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--refresh", action="store_true", help="强制重新拉取 B 站元数据")
    parser.add_argument("--skip-fetch", action="store_true", help="完全离线，只使用缓存")
    args = parser.parse_args()

    registry_rows = read_json(REGISTRY_FILE)["videos"]
    log(f"读取 registry.json：{len(registry_rows)} 条视频")
    main_history = read_history_rows(MAIN_FILE)
    small_history = read_history_rows(SMALL_FILE)
    log(f"读取旧表：主号 {len(main_history)} 条，小号 {len(small_history)} 条")

    bvids = [row["bvid"] for row in registry_rows if row.get("bvid")]
    bvids += [row["bvid"] for row in main_history + small_history if row.get("bvid")]
    bvids = list(dict.fromkeys(bvids))
    meta_cache = refresh_meta_cache(bvids, args.refresh, args.skip_fetch)
    rows, unmatched_main, unmatched_small = build_unified_rows(main_history, small_history, registry_rows, meta_cache)
    log(f"合并后总数：{len(rows)} 条；未匹配：主号 {len(unmatched_main)}，小号 {len(unmatched_small)}")

    backup_file(MAIN_FILE)
    backup_file(SMALL_FILE)
    stamp = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    build_workbook(
        rows,
        unmatched_main,
        "主号历史表",
        f"自动更新：接入 registry.json、B站元数据和公式分析（{stamp}）",
        MAIN_FILE,
        account_filter="主号",
    )
    build_workbook(
        rows,
        unmatched_small,
        "小号历史表",
        f"自动更新：接入 registry.json、B站元数据和公式分析（{stamp}）",
        SMALL_FILE,
        account_filter="小号",
    )
    log("已生成主号表")
    log("已生成小号表")
    log("完成")


if __name__ == "__main__":
    main()
