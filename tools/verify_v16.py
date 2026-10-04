# -*- coding: utf-8 -*-
import os, re
from docx import Document
from docx.oxml.ns import qn

P = r'F:\warma百科\warma-encyclopedia\warma-encyclopedia-v16.docx'
doc = Document(P)

def para_text(el):
    return ''.join(t.text or '' for t in el.findall('.//' + qn('w:t')))

def has_image(el):
    return bool(el.findall('.//' + qn('a:blip'))) or bool(el.findall('.//' + qn('pic:pic')))

children = list(doc.element.body.iterchildren())
img_count = 0
cap_count = 0
ok_caps = 0
seq = []
for i, el in enumerate(children):
    if el.tag == qn('w:p'):
        txt = para_text(el).strip()
        if has_image(el):
            img_count += 1
            nxt = children[i+1] if i+1 < len(children) else None
            ntxt = para_text(nxt).strip() if nxt is not None and nxt.tag == qn('w:p') else ''
            iscap = bool(re.match(r'^图\s*\d+-\d+\s+', ntxt))
            ok_caps += int(iscap)
            seq.append((i, ntxt, iscap))

for p in doc.paragraphs:
    if re.match(r'^图\s*\d+-\d+\s+', p.text.strip()):
        cap_count += 1

img_rels = sum(1 for rel in doc.part.rels.values() if 'image' in rel.reltype)
heads = []
for p in doc.paragraphs:
    if p.style and p.style.name.startswith('Heading'):
        t = p.text.strip()
        if re.match(r'^(1\.[1-6]\s|3\.1\s|3\.2\s|3\.11\s|A\.6\s)', t):
            heads.append(f'[{p.style.name}] {t}')

print('图片元素:', img_count)
print('紧邻图注:', ok_caps)
print('图注文本:', cap_count)
print('图片关系:', img_rels)
print('表格:', len(doc.tables))
print('段落:', len(doc.paragraphs))
print('文件大小: %.2f MB' % (os.path.getsize(P)/1024/1024))
print('\n关键标题:')
for h in heads:
    print(' ', h)
print('\n图片-图注顺序:')
for row in seq:
    print(' ', row)
# 粗查文末孤段
print('\n末尾文本:')
for p in doc.paragraphs[-18:]:
    t=p.text.strip()
    if t: print(' ', t[:80])
