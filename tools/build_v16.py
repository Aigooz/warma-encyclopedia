# -*- coding: utf-8 -*-
"""
build_v16.py — 从 v14 重建可视化版：
1) 修正 v15 图注脱离正文的问题
2) 嵌入全部 12 张图表（含新 09-12）
3) 增加可视化索引
"""
import os
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = r'F:\warma百科\warma-encyclopedia'
SRC  = os.path.join(ROOT, 'warma-encyclopedia-v14.docx')
DST  = os.path.join(ROOT, 'warma-encyclopedia-v16.docx')
CHARTS = os.path.join(ROOT, 'charts')

doc = Document(SRC)
body = doc.element.body

# ---------- helpers ----------
def detach(p):
    p._p.getparent().remove(p._p)
    return p

def set_heading_text(p, text):
    # 保留段落属性，仅替换文本 runs
    for r in list(p._p.findall(qn('w:r'))):
        p._p.remove(r)
    p.add_run(text)

def add_heading(level, text):
    p = doc.add_paragraph()
    detach(p)
    p.style = doc.styles[f'Heading {level}']
    p.add_run(text)
    return p

def add_body(text):
    p = doc.add_paragraph()
    detach(p)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x3D, 0x3D, 0x3D)
    return p

def add_image(filename, width_cm=15.5, caption=''):
    img_path = os.path.join(CHARTS, filename)
    if not os.path.exists(img_path):
        raise FileNotFoundError(img_path)
    p = doc.add_paragraph()
    detach(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(img_path, width=Cm(width_cm))
    cap = doc.add_paragraph()
    detach(cap)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run(caption)
    cr.font.size = Pt(9)
    cr.font.italic = True
    cr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    return p, cap

def insert(elements, anchor):
    for el in elements:
        if hasattr(el, '_p'):
            anchor._p.addprevious(el._p)
        else:
            anchor._p.addprevious(el)

def find_para(pred, required=True, label=''):
    for p in doc.paragraphs:
        if pred(p):
            return p
    if required:
        raise RuntimeError(f'未找到定位段落: {label}')
    return None

def h2_startswith(prefix):
    return lambda p: p.style and p.style.name == 'Heading 2' and p.text.strip().startswith(prefix)

def h4_equals(text):
    return lambda p: p.style and p.style.name == 'Heading 4' and p.text.strip() == text

# ---------- 1.4 数据可视化 ----------
old_14 = find_para(h2_startswith('1.4 标志性口头禅'), required=True, label='旧1.4')
elements = []
elements.append(add_heading(2, '1.4 数据可视化'))

def visual_block(title, body, image, caption, width=15.5):
    els = [add_heading(3, title), add_body(body)]
    img, cap = add_image(image, width_cm=width, caption=caption)
    els.extend([img, cap])
    return els

elements += visual_block(
    '年度投稿趋势',
    '年度柱状图显示，warma 自 2015 年首次投稿以来整体保持长期创作；2019 年以 59 部成为投稿峰值年，2020 年后仍维持每年约 20—30 部的稳定更新节奏。',
    '01_yearly_trend.png', '图 1-1 年度投稿趋势（2015—2026）')
elements += visual_block(
    '视频类型分布',
    '类型环形图按已识别标签统计：游戏实况与翻唱/音乐是两大核心板块，直播录像、爆炸电台、生活杂谈等系列构成了多样化补充；“其他”主要是短日常和较难归类的创作内容。',
    '02_type_donut.png', '图 1-2 视频类型分布（共 318 部）', 14.5)
elements += visual_block(
    '月度投稿热力图',
    '年份×月份热力图展示活跃期与休更期：颜色越深表示当月投稿越多，可以直观看到高产季度、系列更新期和阶段性停更。',
    '03_monthly_heatmap.png', '图 1-3 月度投稿热力图（颜色深浅表示投稿数量）')
elements += visual_block(
    '字幕累计增长',
    '累计曲线显示字幕语料从早期短篇逐步扩展，2019—2020 年增长最快，之后仍随长实况、直播录像和电台内容稳步增加，为语言风格与话题分析提供基础。',
    '04_subtitle_growth.png', '图 1-4 字幕累计增长趋势')
elements += visual_block(
    '主号与小号投稿分布',
    '主号 Warma 承载核心内容；warma养鸽场与怒九摸鱼馆补充直播剪辑、短日常和合作内容。多账号共同构成 warma 的内容生态。',
    '05_account_dist.png', '图 1-5 主号与小号投稿分布', 13.5)
elements += visual_block(
    '投稿间隔分析',
    '投稿间隔散点图显示大部分更新集中在一周以内，也能看到考试、创作周期或旅行带来的空档；整体节奏兼顾高产与恢复期。',
    '06_upload_gaps.png', '图 1-6 投稿间隔分布（虚线为参考线）')
elements += visual_block(
    '字幕量最多的视频',
    'TOP20 字幕量排名以长直播录像、系列实况和大型电台为主；它们保留了更多连续对话与即兴表达，是研究 warma 语言生态的重要样本。',
    '07_top_subtitle.png', '图 1-7 字幕量最多的 20 部视频', 15)
elements += visual_block(
    '「我是warma/沃玛」出现频率',
    '按半年统计的自我介绍频次在早期和特殊节点更高；随着观众熟悉度提升，正式自我介绍逐渐减少，也反映频道从陌生接触到稳定陪伴的关系变化。',
    '08_intro_freq.png', '图 1-8 自我介绍频率变化（按半年）')
insert(elements, old_14)
set_heading_text(old_14, '1.5 标志性口头禅与语言习惯')

old_15 = find_para(h2_startswith('1.5 创作里程碑'), required=True, label='旧1.5')
set_heading_text(old_15, '1.6 创作里程碑与粉丝增长')

# ---------- 1.5 语言章节追加 09/10 ----------
anchor_h4 = find_para(h4_equals('标志性口头禅使用频率TOP20'), required=True, label='口头禅表标题')
lang_els = []
lang_els += visual_block(
    '话题关键词频次',
    '按主题关键词聚类统计，“社交/合作”“日常生活”“美食”“情感”等日常表达出现频率最高；游戏、绘画/创作和音乐紧随其后。这里的“社交/合作”多来自连续对话和合作实况中的“我们/一起/朋友”等表达，不宜直接等同于专题合作视频数量。',
    '09_topics.png', '图 1-9 话题关键词频次 TOP10（按字幕词频统计）')
lang_els += visual_block(
    '口头禅与语气词频率',
    '图表与第 1.5 节统计表使用同一口径：“诶”“哇”“哈哈”“嘿嘿”“天哪”“拜拜”等高频表达构成她即时反应和情绪表达的基础，也让视频保持轻松、真实和强陪伴感。',
    '10_catchphrases.png', '图 1-10 口头禅/语气词使用频率 TOP15（依据 1.5 统计表）')
insert(lang_els, anchor_h4)

# ---------- 第三章：先顺延编号，再插入新 3.1 ----------
for n in range(10, 0, -1):
    p = find_para(h2_startswith(f'3.{n} '), required=False)
    if p is not None:
        old_title = p.text.strip()
        suffix = old_title[len(f'3.{n} '):]
        set_heading_text(p, f'3.{n+1} {suffix}')

anchor31 = find_para(h2_startswith('3.2 游戏实况'), required=True, label='游戏实况')
type_els = []
type_els.append(add_heading(2, '3.1 类型演变与体量对比'))
type_els.append(add_body(
    '本章前两图从时间演变和内容体量两个角度观察 warma 的类型分布：前者关注不同阶段的内容重心变化，后者比较各类型在视频数量与平均字幕量上的差异。'
))
type_els += visual_block(
    '内容类型年度演变',
    '堆叠面积图将相似标签合并为六大类。可以看到，游戏实况长期占据重要位置，音乐/翻唱在早期到中期持续活跃，直播录像、爆炸电台和生活杂谈则在不同阶段形成补充，整体内容结构逐步多元化。',
    '11_type_evolution.png', '图 3-1 内容类型年度演变（堆叠面积图）', 16)
type_els += visual_block(
    '类型×平均字幕量对比',
    '气泡图比较各简化类型的视频数量、平均字幕量和总字幕量：直播录像、爆炸电台通常单部体量更大，游戏实况数量多且累计语料丰富，而短日常和音乐作品则更精炼。',
    '12_type_scatter.png', '图 3-2 内容类型×平均字幕量（气泡大小表示总字幕量）')
insert(type_els, anchor31)

# ---------- 附录 A.6 可视化索引 ----------
end_marks = [p for p in doc.paragraphs if p.text.strip() == '◆ ◆ ◆']
if not end_marks:
    raise RuntimeError('未找到文末装饰段落')
end_anchor = end_marks[-1]
index_els = []
index_els.append(add_heading(2, 'A.6 可视化索引与互动看板'))
index_els.append(add_body('本版共嵌入 12 张静态图表：图 1-1 至图 1-10 分布在第一章“数据可视化”和“标志性口头禅与语言习惯”，图 3-1 至图 3-2 位于第三章“类型演变与体量对比”。所有 PNG 源文件保存在 charts 目录，便于后续按新版数据重新生成。'))
index_els.append(add_body('配套交互看板为 warma-dashboard.html，提供投稿检索、年度/类型/账号/月度/累计增长/投稿间隔等可交互图表。建议将 HTML 与 registry.json 放在同一目录中打开；若浏览器禁止本地 fetch，可使用本地静态服务器运行。'))
index_els.append(add_body('后续更新时，可先更新 registry.json 与 docx 词条，再重新运行图表生成脚本刷新 PNG；本索引会作为可视化模块的固定位置。'))
insert(index_els, end_anchor)

doc.save(DST)

img_rels = 0
d2 = Document(DST)
for rel in d2.part.rels.values():
    if 'image' in rel.reltype:
        img_rels += 1
print(f'✅ 已保存: {DST}')
print(f'图片关系: {img_rels}')
print(f'文件大小: {os.path.getsize(DST)/1024/1024:.2f} MB')
