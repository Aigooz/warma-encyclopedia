from openpyxl import load_workbook
from pathlib import Path

files = [
    Path(r"F:\warma百科\@Warma 相关.xlsx"),
    Path(r"F:\warma百科\@warma养鸽场 相关.xlsx"),
]

for f in files:
    wb = load_workbook(f, data_only=False, read_only=True)
    ws = wb["数据总表"]
    print("\n", f.name)
    for row in ws.iter_rows(min_row=2, max_row=8, max_col=28):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                print(c.coordinate, c.value)
    wb.close()
