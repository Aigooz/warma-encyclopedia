from openpyxl import load_workbook
from pathlib import Path

files = [
    Path(r"F:\warma百科\@Warma 相关.xlsx"),
    Path(r"F:\warma百科\@warma养鸽场 相关.xlsx"),
]

for f in files:
    print("\nFILE", f, "size", f.stat().st_size)
    wb = load_workbook(f, data_only=False, read_only=True)
    print("sheets", wb.sheetnames)
    ws = wb["数据总表"]
    print("dims", ws.max_row, ws.max_column)
    print("headers", [c.value for c in next(ws.iter_rows(min_row=1, max_row=1, max_col=ws.max_column))])
    for sn in ["仪表盘", "年度分析", "类型分析", "时长分段", "数据质量"]:
        if sn in wb.sheetnames:
            s = wb[sn]
            formulas = []
            for row in s.iter_rows(min_row=1, max_row=min(s.max_row, 80), max_col=min(s.max_column, 20)):
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        formulas.append((c.coordinate, c.value))
            print(sn, "dims", s.max_row, s.max_column, "formula_count", len(formulas), "first", formulas[:10])
    wb.close()
