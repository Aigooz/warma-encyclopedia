# -*- coding: utf-8 -*-
"""百科docx增量更新引擎：新增词条 + 更新年表/统计/封面。"""
import copy, re, os, glob
from collections import Counter, OrderedDict
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import RGBColor

def _rgb(c):
    return RGBColor(int(c[0:2], 16), int(c[2:4], 16), int(c[4:6], 16))

C_LABEL = '0984E3'   # 蓝色小标题
C_BODY  = '3D3D3D'   # 正文
C_META  = '53595D'   # 氛围/语言
C_NOTE  = 'B2BEC3'   # 注释/完标记
C_STAR  = 'A29BFE'   # ✦

class DocPatcher:
    def __init__(self, path):
        self.path = path
        self.doc = docx.Document(path)
        self.body = self.doc.element.body
        self._collect_templates()
        self.paras = self.doc.paragraphs

    # ---------- 模板 ----------
    def _collect_templates(self):
        tpls = {}
        for p in self.doc.paragraphs:
            t = p.text.strip()
            if 'tpl_h3' not in tpls and p.style.name == 'Heading 3' and re.match(r'^第\d+部', t):
                tpls['tpl_h3'] = p._p
            elif 'tpl_label' not in tpls and t == '内容摘要：':
                tpls['tpl_label'] = p._p
            elif 'tpl_body' not in tpls and p.style.name == 'Normal' and len(t) > 40 and p.style.name == 'Normal':
                for r in p.runs:
                    if r.font.size and r.font.size.pt == 10 and not r.bold:
                        tpls['tpl_body'] = p._p
                        break
            elif 'tpl_meta' not in tpls and t.startswith('整体氛围：'):
                tpls['tpl_meta'] = p._p
            elif 'tpl_lang' not in tpls and t.startswith('字数') and '个语段' in t:
                tpls['tpl_lang'] = p._p
            elif 'tpl_note' not in tpls and t.startswith('（共') and '行字幕' in t:
                tpls['tpl_note'] = p._p
            elif 'tpl_end' not in tpls and re.match(r'^— 第\d+部完 —$', t):
                tpls['tpl_end'] = p._p
            elif 'tpl_star' not in tpls and t == '✦':
                tpls['tpl_star'] = p._p
            elif 'tpl_list' not in tpls and p.style.name == 'List Paragraph' and t:
                tpls['tpl_list'] = p._p
            elif 'tpl_h2year' not in tpls and p.style.name == 'Heading 2' and re.match(r'^2\.\d+\s*\d{4}年：', t):
                tpls['tpl_h2year'] = p._p
            elif 'tpl_h4year' not in tpls and p.style.name == 'Heading 4' and '年投稿列表' in t:
                tpls['tpl_h4year'] = p._p
        # 字幕正文模板：取字幕全文标记后第一段
        for i, p in enumerate(self.doc.paragraphs):
            if p.text.strip().startswith('字幕全文：'):
                for q in self.doc.paragraphs[i+1:i+4]:
                    if q.text.strip().startswith('（共'):
                        continue
                    if q.style.name == 'Normal' and q.text.strip():
                        tpls['tpl_sub'] = q._p
                        break
                break
        self.tpl = {k: copy.deepcopy(v) for k, v in tpls.items()}
        # 表格模板：词条信息表 + 年份表
        for t in self.doc.tables:
            hdr = [c.text.strip() for c in t.rows[0].cells]
            if 'tbl_info' not in self.tpl and hdr[:2] == ['项目', '内容'] and len(t.rows) <= 11:
                self.tpl['tbl_info'] = copy.deepcopy(t._tbl)
            if 'tbl_year' not in self.tpl and hdr[:4] == ['编号', '日期', '标题', '类型']:
                self.tpl['tbl_year'] = copy.deepcopy(t._tbl)
        assert 'tpl_h3' in self.tpl and 'tbl_info' in self.tpl and 'tbl_year' in self.tpl, '模板定位失败'

    # ---------- 基础构件 ----------
    def _mk_para(self, tpl_key, text, bold=None, size=None, color=None):
        p = copy.deepcopy(self.tpl[tpl_key])
        for r in p.findall(qn('w:r')):
            p.remove(r)
        r = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        if bold:
            b = OxmlElement('w:b'); rpr.append(b)
        if size:
            sz = OxmlElement('w:sz'); sz.set(qn('w:val'), str(int(size * 2))); rpr.append(sz)
            szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(int(size * 2))); rpr.append(szcs)
        if color:
            c = OxmlElement('w:color'); c.set(qn('w:val'), color); rpr.append(c)
        r.append(rpr)
        lines = text.split('\n')
        for i, line in enumerate(lines):
            t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = line
            r.append(t)
            if i < len(lines) - 1:
                r.append(OxmlElement('w:br'))
        p.append(r)
        return p

    def _label_para(self, label, content=''):
        p = copy.deepcopy(self.tpl['tpl_label'])
        runs = p.findall(qn('w:r'))
        for r in runs:
            p.remove(r)
        r1 = OxmlElement('w:r')
        rpr = copy.deepcopy(self.tpl['tpl_label'].findall(qn('w:r'))[0].find(qn('w:rPr')))
        r1.append(rpr)
        t = OxmlElement('w:t'); t.text = label; r1.append(t)
        p.append(r1)
        if content:
            r2 = OxmlElement('w:r')
            r2.append(OxmlElement('w:rPr'))
            t2 = OxmlElement('w:t'); t2.set(qn('xml:space'), 'preserve'); t2.text = content
            r2.append(t2)
            p.append(r2)
        return p

    def _list_para(self, text):
        p = copy.deepcopy(self.tpl['tpl_list'])
        for r in p.findall(qn('w:r')):
            p.remove(r)
        r = OxmlElement('w:r')
        rpr = OxmlElement('w:rPr')
        c = OxmlElement('w:color'); c.set(qn('w:val'), C_BODY); rpr.append(c)
        r.append(rpr)
        t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
        r.append(t)
        p.append(r)
        return p

    def _tbl_info(self, fields):
        tbl = copy.deepcopy(self.tpl['tbl_info'])
        rows = tbl.findall(qn('w:tr'))
        row_tpl = copy.deepcopy(rows[1])
        for r in rows[1:]:
            tbl.remove(r)
        for k, v in fields:
            tr = copy.deepcopy(row_tpl)
            cells = tr.findall(qn('w:tc'))
            self._cell_set(cells[0], k)
            self._cell_set(cells[1], v)
            tbl.append(tr)
        return tbl

    def _cell_set(self, tc, text):
        ps = tc.findall(qn('w:p'))
        p = ps[0]
        runs = p.findall(qn('w:r'))
        if runs:
            keep = runs[0]
            for r in runs[1:]:
                p.remove(r)
            for t in keep.findall(qn('w:t')):
                keep.remove(t)
            t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
            keep.append(t)
        else:
            r = OxmlElement('w:r')
            t = OxmlElement('w:t'); t.text = text
            r.append(t); p.append(r)
        for extra in ps[1:]:
            tc.remove(extra)

    # ---------- 定位 ----------
    def find_h1(self, prefix):
        for p in self.doc.paragraphs:
            if p.style.name == 'Heading 1' and p.text.strip().startswith(prefix):
                return p
        return None

    def find_entry_anchor(self):
        """最后一个词条✦段落（附录H1之前的那个）"""
        h1 = self.find_h1('附录')
        prev = h1._p.getprevious()
        while prev is not None:
            if prev.tag.endswith('}p'):
                txt = ''.join(t.text or '' for t in prev.iter(qn('w:t')))
                if txt.strip() == '✦':
                    return prev
            prev = prev.getprevious()
        raise RuntimeError('未找到插入锚点')

    def find_year_sections(self):
        """{year: {'h2': para, 'table': tbl_element}}"""
        out = {}
        cur_year = None
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        for child in self.body.iterchildren():
            if child.tag.endswith('}p'):
                p = Paragraph(child, self.doc)
                m = re.match(r'^2\.\d+\s*(\d{4})年：(\d+)部投稿$', p.text.strip())
                if p.style.name == 'Heading 2' and m:
                    cur_year = int(m.group(1))
                    out[cur_year] = {'h2': p, 'count': int(m.group(2)), 'table': None}
            elif child.tag.endswith('}tbl') and cur_year is not None:
                t = Table(child, self.doc)
                hdr = [c.text.strip() for c in t.rows[0].cells]
                if hdr[:4] == ['编号', '日期', '标题', '类型'] and out[cur_year]['table'] is None:
                    out[cur_year]['table'] = child
        return out

    # ---------- 词条插入 ----------
    def entry_elements(self, v):
        """v: {no,title,date,type,account,intro_count,sub_lines,sub_chars,bvid,analysis}"""
        a = v['analysis']
        els = []
        els.append(self._mk_para('tpl_h3', f"第{v['no']}部  {v['title']}"))
        els.append(self._tbl_info([
            ('投稿时间', v['date'].replace('-', '/')),
            ('距上次投稿', v.get('gap_note', '首期投稿')),
            ('视频类型', v['type']),
            ('账号', v['account']),
            ('自我介绍次数', f"{v['intro_count']} 次"),
            ('字幕行数', f"{v['sub_lines']} 行"),
            ('字幕字数', f"{v['sub_chars']} 字符"),
            ('视频链接', f"https://www.bilibili.com/video/{v['bvid']}" if v['bvid'] else '—'),
        ]))
        els.append(self._label_para('关键话题：', a['topic_str']))
        if a.get('catch_str'):
            els.append(self._label_para('标志性口头禅：', a['catch_str']))
        els.append(self._label_para('内容摘要：'))
        els.append(self._mk_para('tpl_body', a['summary']))
        els.append(self._label_para('内容亮点：'))
        for h in a['highlights']:
            els.append(self._list_para(h))
        els.append(self._label_para('氛围分析：'))
        els.append(self._mk_para('tpl_meta', '\n'.join(a['atmo'])))
        els.append(self._label_para('语言特征：'))
        els.append(self._mk_para('tpl_lang', a['lang']))
        els.append(self._label_para('字幕全文：'))
        merged = a['merged']
        els.append(self._mk_para('tpl_note', f"（共{v['sub_lines']}行字幕，合并为{len(merged)}段）"))
        for mtext in merged:
            els.append(self._mk_para('tpl_sub', mtext))
        els.append(self._mk_para('tpl_end', f"— 第{v['no']}部完 —"))
        els.append(self._mk_para('tpl_star', '✦'))
        return els

    def append_entries(self, videos):
        cursor = self.find_entry_anchor()
        for v in videos:
            for el in self.entry_elements(v):
                cursor.addnext(el)
                cursor = el

    # ---------- 第二章 ----------
    def _table_row_clone(self, tbl, cells):
        rows = tbl.findall(qn('w:tr'))
        tpl = copy.deepcopy(rows[-1])
        tr = copy.deepcopy(tpl)
        tcs = tr.findall(qn('w:tc'))
        for tc, text in zip(tcs, cells):
            self._cell_set(tc, text)
        tbl.append(tr)
        return tr

    def update_years(self, year_counts):
        """year_counts: {year: [(编号,日期,标题,类型), ...]}"""
        secs = self.find_year_sections()
        for year, rows in sorted(year_counts.items()):
            if year in secs and secs[year]['table'] is not None:
                h2 = secs[year]['h2']
                old_n = secs[year]['count']
                new_n = old_n + len(rows)
                self._rewrite_text(h2._p, re.sub(r'：\d+部投稿$', f'：{new_n}部投稿', h2.text.strip()))
                tbl = secs[year]['table']
                for r in rows:
                    self._table_row_clone(tbl, [str(r[0]), r[1], r[2], r[3]])
            else:
                # 新年份：插入到第三章之前
                idx = max(int(re.match(r'^2\.(\d+)', s['h2'].text.strip()).group(1))
                          for s in secs.values()) + 1
                h2el = copy.deepcopy(self.tpl['tpl_h2year'])
                self._rewrite_text(h2el, f'2.{idx} {year}年：{len(rows)}部投稿')
                h4el = copy.deepcopy(self.tpl['tpl_h4year'])
                self._rewrite_text(h4el, f'{year}年投稿列表')
                tbl = copy.deepcopy(self.tpl['tbl_year'])
                rows_el = tbl.findall(qn('w:tr'))
                for r in rows_el[1:]:
                    tbl.remove(r)
                for r in rows:
                    self._table_row_clone(tbl, [str(r[0]), r[1], r[2], r[3]])
                h1 = self.find_h1('第三章')
                h1._p.addprevious(h2el)
                h1._p.addprevious(h4el)
                h1._p.addprevious(tbl)

    def _rewrite_text(self, p_el, new_text):
        runs = p_el.findall(qn('w:r'))
        if not runs:
            r = OxmlElement('w:r')
            t = OxmlElement('w:t'); t.text = new_text
            r.append(t); p_el.append(r)
            return
        keep = runs[0]
        for r in runs[1:]:
            p_el.remove(r)
        for t in keep.findall(qn('w:t')):
            keep.remove(t)
        for br in keep.findall(qn('w:br')):
            keep.remove(br)
        t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = new_text
        keep.append(t)

    # ---------- 第一章统计表 ----------
    def update_ch1_tables(self, reg):
        from docx.table import Table
        tables = self.doc.tables
        t_year = t_type = None
        for t in tables:
            hdr = [c.text.strip() for c in t.rows[0].cells]
            if hdr[:3] == ['年份', '投稿数', '占比'] and t_year is None:
                t_year = t
            elif hdr[:3] == ['类型', '数量', '占比'] and t_type is None:
                t_type = t
        total = len(reg['videos'])
        ycount = Counter(v['date'][:4] for v in reg['videos'] if v['date'])
        # 年份表：逐行重写
        seen_years = []
        for row in t_year.rows[1:]:
            y = row.cells[0].text.strip()
            seen_years.append(y)
            n = ycount.get(y, 0)
            self._cell_set(row._tr.findall(qn('w:tc'))[0], y)
            self._cell_set(row._tr.findall(qn('w:tc'))[1], str(n))
            self._cell_set(row._tr.findall(qn('w:tc'))[2], f'{n / total * 100:.1f}%' if total else '0%')
        for y in sorted(set(ycount) - set(seen_years)):
            n = ycount[y]
            self._table_row_clone(t_year._tbl, [y, str(n), f'{n / total * 100:.1f}%'])
        # 类型表
        tycount = Counter(v['type'] for v in reg['videos'])
        seen_types = []
        for row in t_type.rows[1:]:
            ty = row.cells[0].text.strip()
            seen_types.append(ty)
            n = tycount.get(ty, 0)
            tcs = row._tr.findall(qn('w:tc'))
            self._cell_set(tcs[1], str(n))
            self._cell_set(tcs[2], f'{n / total * 100:.1f}%' if total else '0%')
        for ty in sorted(set(tycount) - set(seen_types)):
            n = tycount[ty]
            self._table_row_clone(t_type._tbl, [ty, str(n), f'{n / total * 100:.1f}%'])

    # ---------- 附录 ----------
    def update_appendix(self, reg, new_videos):
        from docx.table import Table
        def by_header(sig):
            for t in self.doc.tables:
                hdr = [c.text.strip() for c in t.rows[0].cells]
                if hdr[:len(sig)] == list(sig):
                    yield t
        # A.1 追加
        for t in by_header(['编号', '日期', '标题', '类型', '距上次(天)']):
            for v in new_videos:
                self._table_row_clone(t._tbl, [str(v['no']), v['date'].replace('-', '/'),
                                               v['title'], v['type'], v.get('gap_days', '首期投稿')])
            break
        # A.2 / A.3 / A.4 重建
        vids = [v for v in reg['videos'] if v['date']]
        gaps = []
        for i in range(1, len(vids)):
            from datetime import datetime
            d0 = datetime.strptime(vids[i-1]['date'], '%Y-%m-%d')
            d1 = datetime.strptime(vids[i]['date'], '%Y-%m-%d')
            gaps.append({'no': vids[i]['no'], 'title': vids[i]['title'],
                         'date': vids[i]['date'], 'days': (d1 - d0).days})
        def rebuild(t, data, unit):
            rows = t.rows
            tpl_row = copy.deepcopy(rows[-1]._tr)
            for row in rows[1:]:
                t._tbl.remove(row._tr)
            for rank, item in enumerate(data, 1):
                tr = copy.deepcopy(tpl_row)
                tcs = tr.findall(qn('w:tc'))
                vals = [str(rank), str(item['no']), item['title'],
                        item['date'].replace('-', '/'), f"{item['days']} {unit}"]
                for tc, text in zip(tcs, vals):
                    self._cell_set(tc, text)
                t._tbl.append(tr)
        sig2 = ['排名', '编号', '标题', '日期', '间隔天数']
        gap_tabs = list(by_header(sig2))
        if gap_tabs:
            rebuild(gap_tabs[0], sorted(gaps, key=lambda g: -g['days'])[:10], '天')
        if len(gap_tabs) > 1:
            rebuild(gap_tabs[1], sorted(gaps, key=lambda g: g['days'])[:10], '天')
        for t in by_header(['排名', '编号', '标题', '日期', '字幕行数']):
            rows = t.rows
            tpl_row = copy.deepcopy(rows[-1]._tr)
            for row in rows[1:]:
                t._tbl.remove(row._tr)
            top = sorted(reg['videos'], key=lambda v: -v['sub_lines'])[:10]
            for rank, v in enumerate(top, 1):
                tr = copy.deepcopy(tpl_row)
                tcs = tr.findall(qn('w:tc'))
                vals = [str(rank), str(v['no']), v['title'], (v['date'] or '').replace('-', '/'), f"{v['sub_lines']} 行"]
                for tc, text in zip(tcs, vals):
                    self._cell_set(tc, text)
                t._tbl.append(tr)
            break
        # A.5 追加自我介绍
        for t in by_header(['编号', '标题', '日期', '自我介绍次数', '读音']):
            for v in new_videos:
                if v['intro_count'] > 0:
                    self._table_row_clone(t._tbl, [str(v['no']), v['title'], v['date'].replace('-', '/'),
                                                   f"{v['intro_count']} 次", v.get('reading') or '—'])
            break

    # ---------- 封面 / 统计文本 ----------
    def update_cover(self, reg):
        total = len(reg['videos'])
        main = sum(1 for v in reg['videos'] if '小号' not in (v.get('account') or ''))
        alt = total - main
        lines_total = reg['total_sub_lines']
        dates = sorted(v['date'] for v in reg['videos'] if v['date'])
        span = f'时间跨度：{dates[0]} 至 {dates[-1]}' if dates else None
        repl = [
            (re.compile(r'^收录视频：'), f'收录视频：{total} 部（主号{main}部 + 小号{alt}部）'),
            (re.compile(r'^时间跨度：'), span),
            (re.compile(r'^字幕总行数：'), f'字幕总行数：{lines_total:,} 行'),
            (re.compile(r'^共收录 '), f'共收录 {total} 部视频 | 字幕总行数 {lines_total:,} 行'),
        ]
        for p in self.doc.paragraphs:
            t = p.text.strip()
            for pat, new in repl:
                if new and pat.match(t):
                    self._rewrite_text(p._p, new)
                    break
        h1 = self.find_h1('第十二章')
        if h1:
            m = re.match(r'^(第十二章\s*视频词条详录（全)(\d+)(部）)$', h1.text.strip())
            if m:
                self._rewrite_text(h1._p, f'{m.group(1)}{total}{m.group(3)}')

    # ---------- 保存 ----------
    def save_next_version(self):
        m = re.search(r'-v(\d+)\.docx$', self.path)
        nxt = int(m.group(1)) + 1 if m else 12
        out = self.path.replace(f'-v{m.group(1)}.docx', f'-v{nxt}.docx') if m else self.path
        self.doc.save(out)
        return out


