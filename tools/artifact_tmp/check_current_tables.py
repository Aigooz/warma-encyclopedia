import json
import re
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

base = Path(r'F:\warma百科')
registry = json.loads((base / 'warma-encyclopedia/registry.json').read_text(encoding='utf-8'))
reg_by_bvid = {v['bvid']: v for v in registry['videos']}
reg_by_title = {re.sub(r'\s+', '', v['title']): v for v in registry['videos']}
meta = json.loads((base / 'warma-encyclopedia/tools/bili_meta_cache.json').read_text(encoding='utf-8'))

files = [base / '@Warma 相关.xlsx', base / '@warma养鸽场 相关.xlsx']
for fn in files:
    print('\n===', fn.name, '===')
    wb = load_workbook(fn, data_only=False, read_only=True)
    ws = wb['数据总表']
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row = tuple(row) + (None,) * (28 - len(row))
        title = row[1]
        if not title:
            continue
        link = row[25] or row[11] or row[6] or ''
        bv = None
        m = re.search(r'(BV[0-9A-Za-z]{10})', str(link))
        if m:
            bv = m.group(1)
        rows.append({'title': str(title).strip(), 'bv': bv, 'row': len(rows) + 2, 'link': str(link) if link else None})
    print('data rows', len(rows))
    print('missing bv', sum(1 for r in rows if not r['bv']))
    print('registry matched', sum(1 for r in rows if r['bv'] in reg_by_bvid))
    print('registry title matched', sum(1 for r in rows if re.sub(r'\s+', '', r['title']) in reg_by_title))
    print('registry not matched by bv or title', [r for r in rows if not r['bv'] and re.sub(r'\s+', '', r['title']) not in reg_by_title][:30])
    print('duplicate bvids', [b for b, c in Counter(r['bv'] for r in rows if r['bv']).items() if c > 1][:20])
    wb.close()
