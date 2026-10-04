import json
from pathlib import Path
from openpyxl import load_workbook

base = Path(r'F:\warma百科')
files = [base / '@Warma 相关.xlsx', base / '@warma养鸽场 相关.xlsx']

for fn in files:
    print('\n===', fn.name, '===')
    wb = load_workbook(fn, data_only=False, read_only=False)
    for ws in wb.worksheets:
        print(ws.title, ws.max_row, ws.max_column, 'charts', len(ws._charts), 'tables', len(ws.tables), 'cond', len(ws.conditional_formatting._cf_rules))
        if ws.title in ['数据总表', '年度分析', '类型分析', '时长分段', '仪表盘', '数据质量', '更新日志']:
            hdr = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            print('hdr', hdr)
            for r in range(1, min(ws.max_row, 20) + 1):
                formulas = []
                for c in range(1, ws.max_column + 1):
                    v = ws.cell(r, c).value
                    if isinstance(v, str) and v.startswith('='):
                        formulas.append((ws.cell(r, c).coordinate, v))
                if formulas:
                    print('row', r, formulas)
