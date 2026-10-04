import sys, json, re
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding='utf-8')
base=Path(r'F:\warma百科')
files=[base/'@Warma 相关.xlsx', base/'@warma养鸽场 相关.xlsx']
raw=json.loads((base/'warma-encyclopedia/registry.json').read_text(encoding='utf-8'))
rows=raw['videos']
print('registry',len(rows),raw['docx_version'])
print('registry counts',Counter(v['account'] for v in rows))
for fn in files:
    wb=load_workbook(fn, data_only=False, read_only=False)
    ws=wb.worksheets[0]
    hdr=[c.value for c in ws[1]]
    print('\nWORKBOOK',fn.name)
    print('sheets',wb.sheetnames)
    print('header',hdr,'rows',ws.max_row)
    ids=[]; dates=[]; titles=[]; missing_id=0
    for r in range(2, ws.max_row+1):
        title=ws.cell(r,2).value
        if not title: continue
        link=ws.cell(r,12).value or ws.cell(r,7).value
        bv=None
        if link:
            m=re.search(r'(BV[0-9A-Za-z]{10})',str(link)); bv=m.group(1) if m else None
        if not bv: missing_id+=1
        d=ws.cell(r,4).value
        ids.append(bv); titles.append(title); dates.append(str(d)[:10] if d else None)
    print('tracked rows',len(titles),'missing bv',missing_id,'first/last date',dates[0],dates[-1])
    reg_ids={v['bvid'] for v in rows}
    print('workbook bv found in registry',sum(1 for x in ids if x in reg_ids),'not found',sum(1 for x in ids if x and x not in reg_ids))
    reg_titles={re.sub(r'\s+','',v['title']) for v in rows}
    print('title in registry',sum(1 for t in titles if re.sub(r'\s+','',str(t)) in reg_titles))
    print('sample titles')
    for t in titles[:5]: print(' ',t)
