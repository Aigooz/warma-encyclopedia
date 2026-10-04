from openpyxl import load_workbook
from pathlib import Path

files = [
    Path(r"F:\warma百科\@Warma 相关.xlsx"),
    Path(r"F:\warma百科\@warma养鸽场 相关.xlsx"),
]

for f in files:
    wb = load_workbook(f, data_only=True, read_only=True)
    ws = wb["数据质量"]
    print("\n", f.name)
    for row in ws.iter_rows(min_row=1, max_row=20, max_col=4):
        print([c.value for c in row])
    wb.close()
