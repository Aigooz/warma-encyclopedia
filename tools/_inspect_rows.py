from openpyxl import load_workbook
from pathlib import Path

files = [
    Path(r"F:\warma百科\@Warma 相关.xlsx"),
    Path(r"F:\warma百科\@warma养鸽场 相关.xlsx"),
]

for f in files:
    wb = load_workbook(f, data_only=True, read_only=True)
    ws = wb["数据总表"]
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    nonempty = [r for r in rows[1:] if any(v is not None for v in r)]
    print("\n", f.name, "maxrows", ws.max_row, "nonempty", len(nonempty))
    print("headers", [str(h) for h in headers])
    print("last5")
    for r in nonempty[-5:]:
        print(r[:13])
    missing_bv = [(i+1, r[1], r[2], r[3], r[10], r[11]) for i, r in enumerate(nonempty) if not r[2]]
    missing_av = [(i+1, r[1], r[2], r[3], r[10], r[11]) for i, r in enumerate(nonempty) if not r[3]]
    missing_link = [(i+1, r[1], r[2], r[3], r[10], r[11]) for i, r in enumerate(nonempty) if not r[10]]
    print("missing bv", len(missing_bv), missing_bv[:30])
    print("missing av", len(missing_av), missing_av[:30])
    print("missing link", len(missing_link), missing_link[:30])
    src = {}
    for r in nonempty:
        src[r[11]] = src.get(r[11], 0) + 1
    print("sources", src)
    wb.close()
